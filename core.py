"""Core DNO logic: portfolio parsing, checkup building, DNO template filling.

Framework-free so it can be tested on its own and reused by the web app.
"""
import io
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

MODES = ("sku", "asin", "master_id", "mpn")

TEMPORARY_REASONS = [
    "Excess Inventory", "Listing Issues", "Ongoing Negotiations",
    "Production Delays - Single Item", "Production Delays - Missing Component",
    "Production Delays - Component", "Seasonal Item",
    "SKU Rationalization - Buybox", "SKU Rationalization - Margin",
    "Waiting - Inventory Position", "Waiting - Launch",
]
PERMANENT_REASONS = [
    "Marketplace Regulations", "New Generation", "Partner Termination",
    "Production Discontinued", "Amaczar Legacy DNO", "Listing Deprecated",
]

TEMPLATE_HEADERS = [
    "listing_action", "listing_id", "listing_dno_note", "listing_dno_reason_code",
    "listing_dno_is_temporary", "listing_dno_start_date", "listing_dno_end_date",
    "listing_is_shipable", "listing_sync_avail_inv",
]


# ---------- Helpers ----------
def safe_str(x):
    return str(x).strip() if pd.notna(x) else ""


def normalize_cols(df):
    return {c.strip().lower().replace(" ", "_"): c for c in df.columns}


def find_sheet_like(xls, candidates):
    for s in xls.sheet_names:
        low = s.strip().lower()
        for cand in candidates:
            if cand in low:
                return s
    return None


def find_col(norm_map, candidates):
    for cand in candidates:
        if cand in norm_map:
            return norm_map[cand]
    for k, v in norm_map.items():
        for cand in candidates:
            if cand in k:
                return v
    return None


def validate_date_str(s):
    try:
        datetime.strptime(s, "%Y-%m-%d")
        return True
    except (ValueError, TypeError):
        return False


def parse_values(text):
    """Split comma/newline separated user input into a clean list."""
    return [s.strip() for s in (text or "").replace("\r", "\n").replace(",", "\n").split("\n") if s.strip()]


def _read_sheet(source, sheet):
    try:
        df = pd.read_excel(source, sheet_name=sheet, header=2, dtype=str)
        if df.shape[1] < 2:
            raise ValueError("fallback")
    except Exception:
        df = pd.read_excel(source, sheet_name=sheet, header=0, dtype=str)
    return df.fillna("").map(lambda x: x.strip() if isinstance(x, str) else x)


# ---------- Step 1: portfolio ----------
PROD_REQUIRED = {
    "sku": ["listing_mp_primary_id", "listing_mp_primaryid", "mp_primary_id", "sku", "product_sku", "mp_primary"],
    "listing_id": ["listing_id", "listingid", "mp_listing_id", "mp_list_id"],
    "marketplace": ["listing_mp_name", "listingmpname", "mp_name", "marketplace"],
    "seller": ["listing_mp_seller_name", "listingmp_seller_name", "seller_name", "seller_account",
               "sellername", "seller_id", "sellerid"],
    "master": ["master_id", "masterid", "master"],
}
ASIN_CANDIDATES = ["listing_mp_page_id", "listing_mp_pageid", "mp_page_id", "page_id", "asin", "mp_pageid"]


def load_portfolio(file_like):
    """Parse the portfolio workbook. Returns a picklable dict of state."""
    data = file_like.read() if hasattr(file_like, "read") else file_like
    xls = pd.ExcelFile(io.BytesIO(data))

    product_sheet = find_sheet_like(
        xls, ["product listing", "product listings", "product_listing", "product_listings"])
    if not product_sheet:
        raise ValueError("Product Listings sheet not found. Sheets: " + ", ".join(xls.sheet_names))

    prod_df = _read_sheet(io.BytesIO(data), product_sheet)
    prod_norm = normalize_cols(prod_df)

    cols, missing = {}, []
    for key, cands in PROD_REQUIRED.items():
        cols[key] = find_col(prod_norm, cands)
        if not cols[key]:
            missing.append(key)
    if missing:
        raise ValueError(f"Missing required Product Listings columns: {missing}. "
                         f"Found: {list(prod_df.columns)}")
    cols["asin"] = find_col(prod_norm, ASIN_CANDIDATES)

    master_df, master_cols = None, {}
    master_sheet = find_sheet_like(
        xls, ["master products", "master_products", "master product", "masterproduct"])
    if master_sheet:
        master_df = _read_sheet(io.BytesIO(data), master_sheet)
        mnorm = normalize_cols(master_df)
        master_cols = {
            "mpn": find_col(mnorm, ["mpn", "mpn_code", "mpn_number"]),
            "master_id": find_col(mnorm, ["master_id", "masterid", "master"]),
        }

    marketplaces = sorted({safe_str(x) for x in prod_df[cols["marketplace"]] if safe_str(x)})
    sellers = sorted({safe_str(x) for x in prod_df[cols["seller"]] if safe_str(x)})
    if not marketplaces or not sellers:
        raise ValueError("No marketplaces or sellers detected in Product Listings sheet.")

    return {
        "prod_df": prod_df, "cols": cols,
        "master_df": master_df, "master_cols": master_cols,
        "marketplaces": marketplaces, "sellers": sellers,
    }


def capabilities(state):
    mc = state["master_cols"]
    return {
        "asin": state["cols"]["asin"] is not None,
        "mpn": bool(state["master_df"] is not None and mc.get("mpn") and mc.get("master_id")),
    }


# ---------- Step 2: checkup ----------
def build_checkup(state, mode, values, scope_all, marketplaces, sellers):
    """Return a list of row dicts. One row per (input value, listing_id)."""
    if mode not in MODES:
        raise ValueError(f"Unknown mode: {mode}")
    cols, prod = state["cols"], state["prod_df"]
    caps = capabilities(state)
    if mode == "asin" and not caps["asin"]:
        raise ValueError("ASIN / listing_mp_page_id column not found in Product Listings; ASIN mode unavailable.")
    if mode == "mpn" and not caps["mpn"]:
        raise ValueError("Master Products sheet or its 'mpn'/'master_id' columns not found; MPN mode unavailable.")
    if not scope_all and (not marketplaces or not sellers):
        raise ValueError("Select at least one marketplace and one seller for Specific Listing IDs.")

    def in_scope(df):
        if scope_all:
            return df
        return df[df[cols["marketplace"]].isin(marketplaces) & df[cols["seller"]].isin(sellers)]

    def rows_for(base, key, empty_sku=""):
        """key: dict of leading columns identifying the input value."""
        base = in_scope(base)
        if base.empty:
            return [{**key, "listing_id": "", "sku": empty_sku, "marketplace": "", "seller_account": "", "count": 0}]
        out = []
        for (mp, seller), sub in base.groupby([cols["marketplace"], cols["seller"]], sort=True):
            lids = sub[cols["listing_id"]].replace("", pd.NA).dropna().astype(str).unique().tolist()
            for lid in lids:
                rep = sub[sub[cols["listing_id"]] == lid].iloc[0]
                out.append({**key, "listing_id": safe_str(lid), "sku": safe_str(rep[cols["sku"]]),
                            "marketplace": safe_str(mp), "seller_account": safe_str(seller),
                            "count": len(lids)})
        return out

    if mode == "mpn" and not values:
        mc, mdf = state["master_cols"], state["master_df"]
        values = mdf[mc["mpn"]].replace("", pd.NA).dropna().astype(str).unique().tolist()
    if not values:
        raise ValueError("Enter at least one value.")

    rows = []
    for v in values:
        if mode == "sku":
            rows += rows_for(prod[prod[cols["sku"]] == v], {}, empty_sku=v)
        elif mode == "asin":
            rows += rows_for(prod[prod[cols["asin"]] == v], {"asin": v})
        elif mode == "master_id":
            rows += rows_for(prod[prod[cols["master"]] == v], {"master_id": v})
        else:  # mpn
            mc, mdf = state["master_cols"], state["master_df"]
            mids = mdf[mdf[mc["mpn"]] == v][mc["master_id"]].replace("", pd.NA).dropna().astype(str).unique().tolist()
            if not mids:
                rows.append({"mpn": v, "master_id": "", "listing_id": "", "sku": "",
                             "marketplace": "", "seller_account": "", "count": 0})
                continue
            for mid in mids:
                rows += rows_for(prod[prod[cols["master"]] == mid], {"mpn": v, "master_id": mid})

    return rows


def checkup_to_xlsx(rows):
    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, index=False, sheet_name="Checkup")
    return buf.getvalue()


# ---------- Step 3: DNO file ----------
def generate_dno(template_path, listing_ids, cfg):
    """Fill the formatted template. cfg keys: temporary(bool), reason, start, end,
    shippable, sync_inv, note. Returns xlsx bytes."""
    start, end = cfg.get("start", "").strip(), cfg.get("end", "").strip()
    if not validate_date_str(start):
        raise ValueError("Invalid start date. Use YYYY-MM-DD.")
    if cfg["temporary"] and not validate_date_str(end):
        raise ValueError("Invalid end date for Temporary DNO. Use YYYY-MM-DD.")
    allowed = TEMPORARY_REASONS if cfg["temporary"] else PERMANENT_REASONS
    if cfg.get("reason") not in allowed:
        raise ValueError("Reason code is not valid for the selected DNO type.")
    if not listing_ids:
        raise ValueError("No listing IDs found in Checkup.")

    wb = load_workbook(template_path)
    if "Product Listings" not in wb.sheetnames:
        raise ValueError("Template must contain a 'Product Listings' sheet.")
    ws = wb["Product Listings"]

    header_row, first_data_row = 3, 4
    header_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is not None:
            header_map[str(val).strip().lower().replace(" ", "_")] = col
    missing = [h for h in TEMPLATE_HEADERS if h not in header_map]
    if missing:
        raise ValueError(f"Template 'Product Listings' sheet missing required columns: {missing}")

    for r in range(first_data_row, ws.max_row + 1):
        for h in TEMPLATE_HEADERS:
            ws.cell(row=r, column=header_map[h]).value = None

    for i, lid in enumerate(listing_ids):
        r = first_data_row + i
        values = {
            "listing_action": "Update Existing",
            "listing_id": lid,
            "listing_dno_note": cfg.get("note", "").strip(),
            "listing_dno_reason_code": cfg["reason"],
            "listing_dno_is_temporary": "Y" if cfg["temporary"] else "N",
            "listing_dno_start_date": start,
            "listing_dno_end_date": end if cfg["temporary"] else "",
            "listing_is_shipable": cfg.get("shippable", "") or "",
            "listing_sync_avail_inv": cfg.get("sync_inv", "") or "",
        }
        for h, v in values.items():
            ws.cell(row=r, column=header_map[h]).value = v

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
