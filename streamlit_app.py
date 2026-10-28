"""Smart DNO Automation - Streamlit version (uses core.py).

Run:  streamlit run streamlit_app.py
"""
import io
from datetime import datetime

import pandas as pd
import streamlit as st

import core

st.set_page_config(page_title="Smart DNO Automation", layout="wide")
st.title("Smart DNO Automation")


@st.cache_resource
def template_store():
    """Shared across sessions while the app process is alive."""
    return {}


store = template_store()
ss = st.session_state

MODE_LABELS = {"sku": "By SKU", "asin": "By ASIN", "master_id": "By Master ID", "mpn": "By MPN"}


# ---------- Step 0: template ----------
st.header("Step 0 - DNO template")
if "bytes" in store:
    st.success(f"Template saved: {store['name']} (reused until the app restarts)")
tpl = st.file_uploader("Upload / replace formatted DNO template (.xlsx)", type=["xlsx"], key="tpl")
if tpl is not None and store.get("name") != tpl.name + str(tpl.size):
    data = tpl.getvalue()
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        if "Product Listings" not in wb.sheetnames:
            raise ValueError("Template must contain a 'Product Listings' sheet.")
        store["bytes"], store["name"] = data, tpl.name + str(tpl.size)
        st.rerun()
    except Exception as e:
        st.error(f"Invalid template: {e}")

# ---------- Step 1: portfolio ----------
st.header("Step 1 - Portfolio file")
pf = st.file_uploader("Portfolio (.xlsx with Product Listings & optional Master Products)",
                      type=["xlsx", "xlsm"], key="pf")
if pf is not None:
    key = (pf.name, pf.size)
    if ss.get("pf_key") != key:
        try:
            with st.spinner("Reading portfolio..."):
                ss["portfolio"] = core.load_portfolio(pf.getvalue())
            ss["pf_key"] = key
            ss.pop("checkup", None)
            ss.pop("dno_file", None)
        except Exception as e:
            ss.pop("portfolio", None)
            ss.pop("pf_key", None)
            st.error(str(e))

state = ss.get("portfolio")
if not state:
    st.info("Upload a portfolio file to continue.")
    st.stop()

caps = core.capabilities(state)
st.caption(f"{len(state['prod_df'])} listing rows - {len(state['marketplaces'])} marketplace(s) - "
           f"{len(state['sellers'])} seller(s)")

# ---------- Step 2: checkup ----------
st.header("Step 2 - Checkup")
modes = [m for m in core.MODES if caps.get(m, True)]
mode = st.radio("Mode", modes, format_func=MODE_LABELS.get, horizontal=True)
brand = st.text_input("Brand (used in file names)")
scope = st.radio("Scope", ["All Listing IDs", "Specific Listing IDs"], horizontal=True)
mps, sellers = [], []
if scope == "Specific Listing IDs":
    c1, c2 = st.columns(2)
    mps = c1.multiselect("Marketplace(s)", state["marketplaces"])
    sellers = c2.multiselect("Seller(s)", state["sellers"])
hint = " - leave blank to use all from Master Products" if mode == "mpn" else ""
values = st.text_area(f"{MODE_LABELS[mode][3:]}s (comma / newline separated){hint}", height=120)

if st.button("Create Checkup", type="primary"):
    ss.pop("dno_file", None)
    if not brand.strip():
        st.error("Brand is required (for final DNO file name).")
    else:
        try:
            rows = core.build_checkup(state, mode, core.parse_values(values),
                                      scope == "All Listing IDs", mps, sellers)
            ss["checkup"] = rows
            ss["brand"] = brand.strip()
        except ValueError as e:
            ss.pop("checkup", None)
            st.error(str(e))

rows = ss.get("checkup")
if not rows:
    st.stop()

found = [r for r in rows if r["listing_id"]]
listing_ids = list(dict.fromkeys(r["listing_id"] for r in found))
st.write(f"**{len(listing_ids)}** unique listing ID(s); {len(rows) - len(found)} input(s) with no match.")
st.dataframe(pd.DataFrame(rows), use_container_width=True, height=300)
safe_brand = "".join(c for c in ss["brand"] if c.isalnum() or c in ("_", "-", " ")).strip() or "Brand"
st.download_button(
    "Download Checkup (.xlsx)", core.checkup_to_xlsx(rows),
    file_name=f"Checkup_{safe_brand}_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
if not listing_ids:
    st.stop()

# ---------- Step 3: DNO ----------
st.header("Step 3 - DNO configuration")
if "bytes" not in store:
    st.warning("Upload the DNO template in Step 0 first.")
    st.stop()

dno_type = st.radio("DNO Type", ["Temporary", "Permanent"], horizontal=True)
temp = dno_type == "Temporary"
reason = st.selectbox("Reason Code", core.TEMPORARY_REASONS if temp else core.PERMANENT_REASONS)
c1, c2 = st.columns(2)
start_d = c1.date_input("Start Date", value=None, format="YYYY-MM-DD")
end_d = c2.date_input("End Date", value=None, format="YYYY-MM-DD") if temp else None
start = start_d.isoformat() if start_d else ""
end = end_d.isoformat() if end_d else ""
c3, c4 = st.columns(2)
ship = c3.selectbox("Shippable", ["", "Y", "N"])
sync = c4.selectbox("Sync Avail Inv", ["", "Y", "N"])
note = st.text_input("DNO Note (optional)")

if st.button("Confirm & Generate DNO Upload", type="primary"):
    cfg = {"temporary": temp, "reason": reason, "start": start, "end": end,
           "shippable": ship, "sync_inv": sync, "note": note}
    try:
        data = core.generate_dno(io.BytesIO(store["bytes"]), listing_ids, cfg)
        ss["dno_file"] = (data, f"{safe_brand} DNO {dno_type} {datetime.now():%d-%m-%Y}.xlsx")
    except ValueError as e:
        ss.pop("dno_file", None)
        st.error(str(e))

if ss.get("dno_file"):
    data, name = ss["dno_file"]
    st.success(f"DNO Upload file generated. Rows written: {len(listing_ids)}")
    st.download_button("Download DNO file", data, file_name=name,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
